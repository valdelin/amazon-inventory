#!/usr/bin/env python3
"""Auditoria de gaps do amazon-inventory (SP-API).

Cruza os caches (cache_mensal_*/cache_historico_*) com os relatórios
(relatorio_mensal_*.xlsx) para responder:

- continuidade: existe `relatorio_mensal_*.xlsx` para todos os meses entre o
  primeiro registrado e o mês corrente?
- cobertura de cache: todo mês com relatório tem cache correspondente (e vice-versa)?
- consistência: os totais do XLSX (Vendas Mensais) batem com os recomputados
  do cache (mesmo código que gerou o relatório)?
- janela: o cache cobre o mês inteiro (ou "1º do mês → hoje" quando é o mês atual)?
- vazios semanais: se o mês corrente tem uma janela antiga (ex.: até dia X),
  quantos dias ficaram de fora até hoje?

Não toca na rede: só lê cache/ e *.xlsx. Saída em texto simples.
"""

import glob
import json
import re
import sys
from datetime import date

from openpyxl import load_workbook
from inventario import CACHE_DIR, parse_flat_file, summarise_orders

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho",
         "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
MESES_INDEX = {nome: i for i, nome in enumerate(MESES)}


def month_key(year, month):
    return f"{year:04d}-{month:02d}"


def month_of_filename(path):
    nome = path.rsplit("/", 1)[-1]
    for y in range(2020, 2032):
        for m in range(1, 13):
            label = f"{MESES[m-1]}_{y}"
            if label in nome:
                return month_key(y, m)
    return None


MONTH_RE = re.compile(r"(\d{4})(\d{2})")


def cache_month(path):
    m = MONTH_RE.search(path.rsplit("/", 1)[-1])
    return month_key(int(m.group(1)), int(m.group(2))) if m else None


def load_summary_from_cache(path):
    payload = json.loads(open(path, encoding="utf-8").read())
    if "rows" in payload:
        header, rows = payload["header"], payload["rows"]
        typo = "historico"
    else:
        header, rows = parse_flat_file(payload["pedidos"])
        typo = "mensal"
    summary = summarise_orders(header, rows)
    return typo, payload.get("start"), payload.get("end"), summary


def xlsx_summary(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Vendas Mensais"]
    values = dict()
    for row in ws.iter_rows(values_only=True):
        label = (row[0] or "").strip()
        if label in ("Pedidos (mês)", "Unidades (mês)", "Total (mês)"):
            values[label] = row[1]
    wb.close()
    window = ws["A1"].value or ""
    return window, values


def main():
    reports = sorted(glob.glob("relatorio_mensal_*.xlsx"))
    caches = sorted(glob.glob(str(CACHE_DIR / "cache_*_*.json")))

    rep_months = {}
    for p in reports:
        rep_months.setdefault(month_of_filename(p), []).append(p)
    cache_by_month = {}
    for c in caches:
        cache_by_month.setdefault(cache_month(c), []).append(c)

    if not rep_months:
        print("Nenhum relatorio encontrado.")
        return 1

    first = min(rep_months)
    today = date.today()
    last = month_key(today.year, today.month)

    print(f"=== Continuidade: {first} -> {last} ===")
    missing_report = []
    missing_cache = []
    for y in range(int(first[:4]), int(last.split("-")[0]) + 1):
        for m in range(1, 13):
            key = month_key(y, m)
            if key < first or key > last:
                continue
            flags = []
            if key not in rep_months:
                missing_report.append(key)
            if key not in cache_by_month:
                missing_cache.append(key)
            if key not in rep_months or key not in cache_by_month:
                continue
            rep = rep_months[key][0]
            for cache in cache_by_month[key]:
                typo, cstart, cend, summary = load_summary_from_cache(cache)
                window, vals = xlsx_summary(rep)
                x_ped = vals.get("Pedidos (mês)")
                x_uni = vals.get("Unidades (mês)")
                x_tot = vals.get("Total (mês)")
                ok_ped = "OK" if x_ped == summary["pedidos"] else f"XLSX={x_ped} vs cache={summary['pedidos']}"
                ok_uni = "OK" if x_uni == summary["unidades"] else f"XLSX={x_uni} vs cache={summary['unidades']}"
                ok_tot = "OK" if abs(float(x_tot or 0) - summary["total"]) < 0.01 else \
                    f"XLSX={x_tot} vs cache={summary['total']}"
                print(f"{key} | {typo:<9} | janela {cstart} -> {cend}")
                print(f"    pedidos/units/total: {ok_ped} | {ok_uni} | {ok_tot}")
                print(f"    (xlsx janela: {window})")

    if missing_report:
        print("MESES SEM RELATÓRIO:", ", ".join(missing_report))
    if missing_cache:
        print("MESES SEM CACHE:", ", ".join(missing_cache))

    print("=== Janela do mês corrente ===")
    for cache in cache_by_month.get(last, []):
        typo, cstart, cend, _ = load_summary_from_cache(cache)
        d_end = date.fromisoformat(cend)
        dias_fora = (today - d_end).days
        print(f"{last}: cache {typo} cobre até {cend} -> {dias_fora} dia(s) fora até hoje")
    if last not in cache_by_month:
        print(f"{last}: sem cache/relatório para o mês corrente até o dia {today}")

    return 0


if __name__ == "__main__":
    sys.exit(main())